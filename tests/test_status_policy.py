"""One status policy, read the same way by every layer.

The queue, the review card, the workbook's Source Info and Runs sheets and the
runs overview must never disagree about whether a document (or one logical run
of a packet) can be used without a person. Each asks ``core.review``; these
tests hold every layer to that answer, on single documents and on packets.
"""

from __future__ import annotations

import importlib
import itertools
from datetime import date
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from core.export import to_bytes
from core.reconcile import reconcile
from core.review import blocks_trust, canonical_run_status, canonical_status
from core.runs import runs_overview
from core.schema import (
    Claim,
    ClaimStatus,
    DocumentStatus,
    Finding,
    FindingScope,
    LineOfBusiness,
    LogicalRun,
    LossRunDocument,
    ReconciliationResult,
    Severity,
)


def _finding(rule: str, category: str, severity: Severity, run_id: str | None = None,
             condition: str | None = None) -> Finding:
    return Finding(
        rule_id=rule, scope=FindingScope.DOCUMENT, subject="document",
        condition=condition or f"{rule}:{run_id}", category=category,
        severity=severity, message=f"{rule} raised", run_id=run_id,
    )


EXTRACTION_WARN = ("R-15", "extraction", Severity.WARN)
EXTRACTION_ERROR = ("R-20", "extraction", Severity.ERROR)
FINANCIAL_ERROR = ("R-04", "financial", Severity.ERROR)
UNDERWRITING_WARN = ("R-13", "underwriting", Severity.WARN)
UNDERWRITING_INFO = ("R-14", "underwriting", Severity.INFO)
KINDS = [EXTRACTION_WARN, EXTRACTION_ERROR, FINANCIAL_ERROR, UNDERWRITING_WARN, UNDERWRITING_INFO]


def _result(findings: list[Finding], run_status: dict | None = None) -> ReconciliationResult:
    engine = (DocumentStatus.NEEDS_REVIEW
              if any(f.severity is Severity.ERROR for f in findings) else DocumentStatus.CLEAN)
    return ReconciliationResult(status=engine, findings=findings, run_status=run_status or {})


def _document(runs: list[LogicalRun] | None = None) -> LossRunDocument:
    claims = [
        Claim(claim_number="A-1", date_of_loss=date(2023, 3, 1), claim_status=ClaimStatus.CLOSED,
              incurred_total=Decimal("10"), source_page=1),
        Claim(claim_number="B-1", date_of_loss=date(2023, 4, 1), claim_status=ClaimStatus.CLOSED,
              incurred_total=Decimal("20"), source_page=2),
    ]
    return LossRunDocument(source_filename="synthetic.pdf", file_sha256="0" * 64, page_count=2,
                           claims=claims, runs=runs or [])


def _packet_runs() -> list[LogicalRun]:
    return [LogicalRun(run_id="run-1", pages=[1], confidence="printed"),
            LogicalRun(run_id="run-2", pages=[2], confidence="printed")]


def _source_info(document: LossRunDocument, result: ReconciliationResult | None) -> dict:
    book = load_workbook(BytesIO(to_bytes(document, result)))
    return {row[0].value: row[1].value for row in book["Source Info"].iter_rows()}


def _runs_sheet(document: LossRunDocument, result: ReconciliationResult) -> dict[str, str]:
    sheet = load_workbook(BytesIO(to_bytes(document, result)))["Runs"]
    header = [cell.value for cell in sheet[1]]
    at = header.index("Status")
    return {row[0].value: row[at].value for row in sheet.iter_rows(min_row=2)}


@pytest.fixture(scope="module")
def app():
    return importlib.import_module("app")


def test_a_warn_that_says_the_document_was_not_read_cleanly_blocks_trust():
    result = _result([_finding(*EXTRACTION_WARN)])
    assert result.status is DocumentStatus.CLEAN            # the engine's narrower answer
    assert canonical_status(result) is DocumentStatus.NEEDS_REVIEW


def test_an_underwriting_observation_never_blocks_trust():
    result = _result([_finding(*UNDERWRITING_WARN), _finding(*UNDERWRITING_INFO)])
    assert canonical_status(result) is DocumentStatus.CLEAN


def test_the_policy_fails_closed_on_what_it_cannot_see():
    assert canonical_status(None) is DocumentStatus.NEEDS_REVIEW
    assert canonical_status(_result([]), needs_mapping=True) is DocumentStatus.NEEDS_REVIEW
    assert canonical_run_status(_result([], {"run-1": DocumentStatus.CLEAN}), "run-9") \
        is DocumentStatus.NEEDS_REVIEW
    unclean_run = _result([], {"run-1": DocumentStatus.CLEAN, "run-2": DocumentStatus.NEEDS_REVIEW})
    assert canonical_status(unclean_run) is DocumentStatus.NEEDS_REVIEW


@pytest.mark.parametrize("combo", [
    combo for size in range(0, 3) for combo in itertools.combinations(KINDS, size)
], ids=lambda combo: "+".join(kind[0] for kind in combo) or "none")
def test_app_and_workbook_never_disagree_about_a_document(app, combo):
    findings = [_finding(*kind, condition=f"{kind[0]}:{i}") for i, kind in enumerate(combo)]
    result = _result(findings)
    expected = canonical_status(result)
    assert expected is (DocumentStatus.NEEDS_REVIEW if any(map(blocks_trust, findings))
                        else DocumentStatus.CLEAN)

    queue = app._status_of(SimpleNamespace(needs_mapping=False, reconciliation=result))
    workbook = _source_info(_document(), result)["Reconciliation status"]
    pill = app._status_pill(SimpleNamespace(needs_mapping=False, reconciliation=result))
    if expected is DocumentStatus.CLEAN:
        assert (queue, workbook) == ("clean", "Reconciled")
        assert "Ready" in pill
    else:
        assert (queue, workbook) == ("needs_review", "Needs review")
        assert "Ready" not in pill


def test_a_workbook_without_a_reconciliation_is_not_reconciled():
    assert _source_info(_document(), None)["Reconciliation status"] == "Needs review"


def test_a_packet_where_one_run_reconciles_and_the_other_does_not(app):
    """Run 2 carries an extraction WARN only -- the engine calls it CLEAN, the
    policy does not. Every layer shows run 1 clean and run 2 and the packet not."""
    runs = _packet_runs()
    document = _document(runs)
    result = _result(
        [_finding(*EXTRACTION_WARN, run_id="run-2"), _finding(*UNDERWRITING_WARN, run_id="run-1")],
        {"run-1": DocumentStatus.CLEAN, "run-2": DocumentStatus.CLEAN},
    )
    assert canonical_run_status(result, "run-1") is DocumentStatus.CLEAN
    assert canonical_run_status(result, "run-2") is DocumentStatus.NEEDS_REVIEW
    assert canonical_status(result) is DocumentStatus.NEEDS_REVIEW

    assert _runs_sheet(document, result) == {"run-1": "CLEAN", "run-2": "NEEDS_REVIEW"}
    overview = {row["Run"]: row["Status"] for row in runs_overview(document, result)}
    assert overview == {"run-1": "CLEAN", "run-2": "NEEDS_REVIEW"}
    assert _source_info(document, result)["Reconciliation status"] == "Needs review"
    assert app._status_of(SimpleNamespace(needs_mapping=False, reconciliation=result)) \
        == "needs_review"


def test_a_finding_raised_about_the_whole_packet_applies_to_every_run():
    result = _result([_finding(*EXTRACTION_WARN)],
                     {"run-1": DocumentStatus.CLEAN, "run-2": DocumentStatus.CLEAN})
    assert {canonical_run_status(result, run) for run in ("run-1", "run-2")} \
        == {DocumentStatus.NEEDS_REVIEW}


def _row_gap_document(seen: int) -> LossRunDocument:
    claims = [Claim(claim_number=f"C-{n}", date_of_loss=date(2023, 1, 10 + n),
                    claim_status=ClaimStatus.CLOSED, paid_total=Decimal("10"),
                    reserve_total=Decimal("0"), incurred_total=Decimal("10"),
                    source_page=1)
              for n in range(2)]
    return LossRunDocument(source_filename="s.pdf", file_sha256="0" * 64, page_count=1,
                           valuation_date=date(2023, 12, 31), claims=claims,
                           rows_seen_per_page={1: seen}, processed_pages=[1])


def test_a_row_count_gap_blocks_trust_without_changing_its_severity():
    """Codex P1 on 2e62e69: R-19 says a claim row was seen and not read.

    It stays the spec's WARN, but a document missing a claim row is not
    reconciled, so every canonical-status consumer reads NEEDS_REVIEW.
    """
    from core.review import trust_class

    result = reconcile(_row_gap_document(seen=3))
    gap = [f for f in result.findings if f.rule_id == "R-19"]
    assert gap and all(f.severity is Severity.WARN for f in gap)
    assert all(blocks_trust(f) for f in gap)
    assert canonical_status(result) is DocumentStatus.NEEDS_REVIEW
    assert trust_class(result) == "unresolved"


def test_without_a_row_count_gap_the_same_document_is_clean():
    result = reconcile(_row_gap_document(seen=2))
    assert not [f for f in result.findings if f.rule_id == "R-19"]
    assert canonical_status(result) is DocumentStatus.CLEAN


def test_a_run_awaiting_its_column_mapping_is_not_clean():
    result = ReconciliationResult(status=DocumentStatus.CLEAN,
                                  run_status={"run-1": DocumentStatus.CLEAN})
    assert canonical_run_status(result, "run-1") is DocumentStatus.CLEAN
    assert canonical_run_status(result, "run-1", needs_mapping=True) \
        is DocumentStatus.NEEDS_REVIEW


def _packet_document(*, second_printed: int | None = 3) -> LossRunDocument:
    """A packet whose page-1 facts differ from its runs' own.

    The runs agree on insured, policy number, period, line and valuation, and
    disagree on carrier. The document-level fields are page 1's, so any Source
    Info row still showing them would be answering for the whole packet with
    one run's facts.
    """
    runs = [
        LogicalRun(run_id="run-1", pages=[1], confidence="printed",
                   carrier="Alpha Mutual", named_insured="Shared Insured",
                   policy_number="POL-9", policy_period_start=date(2022, 1, 1),
                   policy_period_end=date(2022, 12, 31),
                   line_of_business=LineOfBusiness.WC,
                   valuation_date=date(2023, 6, 30), printed_claim_count=2),
        LogicalRun(run_id="run-2", pages=[2], confidence="printed",
                   carrier="Beta Casualty", named_insured="Shared Insured",
                   policy_number="POL-9", policy_period_start=date(2022, 1, 1),
                   policy_period_end=date(2022, 12, 31),
                   line_of_business=LineOfBusiness.WC,
                   valuation_date=date(2023, 6, 30),
                   printed_claim_count=second_printed),
    ]
    return LossRunDocument(
        source_filename="packet.pdf", file_sha256="0" * 64, page_count=2,
        carrier="Page One Mutual", named_insured="Page One Insured",
        policy_number="PAGE-1", policy_period_start=date(1999, 1, 1),
        policy_period_end=date(1999, 12, 31), line_of_business=LineOfBusiness.GL,
        valuation_date=date(1999, 6, 30), printed_claim_count=1,
        claims=_document().claims, runs=runs)


def test_a_packet_source_info_shows_a_field_only_when_every_run_agrees():
    info = _source_info(_packet_document(), None)
    assert info["Carrier"] == "differs by run — see Runs sheet"
    assert info["Named insured"] == "Shared Insured"
    assert info["Policy number"] == "POL-9"
    assert info["Policy period"] == "2022-01-01 to 2022-12-31"
    assert info["Line of business"] == "WC"
    assert info["Valuation date"].date() == date(2023, 6, 30)


def test_a_packet_source_info_sums_each_runs_printed_claim_count():
    info = _source_info(_packet_document(second_printed=3), None)
    assert info["Claim count printed on document"] == 5


def test_a_packet_source_info_says_not_printed_when_a_run_has_no_count():
    info = _source_info(_packet_document(second_printed=None), None)
    assert info["Claim count printed on document"] == "not printed"


def test_a_single_report_source_info_still_shows_its_own_facts():
    document = _document().model_copy(update={"carrier": "Solo Mutual"})
    assert _source_info(document, None)["Carrier"] == "Solo Mutual"


def test_every_summary_of_a_row_count_gap_agrees_with_the_policy():
    """Codex P2 on 67e7624: the headline and the card follow blocks_trust.

    R-19 keeps the spec's WARN severity and is categorised ``extraction`` --
    rows seen on a page and not read are a missing claim -- so it blocks trust
    as a reading problem and no summary may call such a document reconciled.
    """
    from core.review import review_bucket, summarise_review

    result = reconcile(_row_gap_document(seen=3))
    summary = summarise_review(result.findings)
    assert summary.headline() == "not read cleanly"
    gap = [f for f in result.findings if f.rule_id == "R-19"]
    assert all(f.severity is Severity.WARN for f in gap)
    assert all(f.category.value == "extraction" for f in gap)
    assert all(review_bucket(f) == "extraction" for f in gap)
    assert all(f in summary.extraction.findings for f in gap)
    assert not any(f in summary.underwriting.findings for f in gap)
