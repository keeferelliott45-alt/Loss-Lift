"""Red-team attack tests for the XLSX export package.

Every test asserts a *security property* of the workbook bytes, not an
implementation detail. All data is synthetic and built in code.
"""

from __future__ import annotations

import datetime
import io
import re
import zipfile
from decimal import Decimal

import openpyxl
import pytest

from core.account import build_account
from core.export import (
    account_to_bytes,
    build_account_workbook,
    build_workbook,
    suggested_filename,
    to_bytes,
)
from core.schema import (
    REDACTED_FIELDS,
    Claim,
    ClaimStatus,
    DocumentStatus,
    Finding,
    FindingCategory,
    FindingScope,
    LossRunDocument,
    Resolution,
    ReviewAction,
    ReviewLog,
    Severity,
    finding_key,
)

# ---------------------------------------------------------------------------
# Synthetic builders
# ---------------------------------------------------------------------------

FORMULA_PAYLOADS = [
    "=cmd|'/c calc'!A0",
    '=HYPERLINK("http://x","y")',
    "+1+1",
    "-2+3",
    "@SUM(1)",
    "\t=1",
    "\r=1",
    "\n=1",
    " =1",
]

SHA = "ab" * 32


def _claim(**overrides) -> Claim:
    base = dict(
        claim_number="CLM-0001",
        date_of_loss=datetime.date(2023, 5, 4),
        claimant_name="Jane Synthetic",
        loss_description="Slip in the synthetic lobby",
        cause_of_loss="Water",
        claim_status=ClaimStatus.OPEN,
        incurred_total=Decimal("100.00"),
        source_page=1,
    )
    base.update(overrides)
    return Claim(**base)


def _finding(message: str = "synthetic finding") -> Finding:
    return Finding(
        rule_id="R-01",
        severity=Severity.ERROR,
        message=message,
        scope=FindingScope.CLAIM,
        category=FindingCategory.FINANCIAL,
        subject="paid_total",
        condition="paid + reserve == incurred",
        claim_number="CLM-0001",
        field="paid_total",
    )


def _resolution(**overrides) -> Resolution:
    f = _finding()
    base = dict(
        key=finding_key(f),
        action=ReviewAction.CORRECTED,
        reviewer="synthetic-reviewer",
        at=datetime.datetime(2024, 1, 2, 12, 0, 0),
        note="reviewer note",
        rule_id=f.rule_id,
        severity=f.severity.value,
        message=f.message,
        claim_number=f.claim_number,
        field=f.field,
        row_id="row-1",
        where="Claim Detail",
        status_before=DocumentStatus.NEEDS_REVIEW.value,
        status_after=DocumentStatus.CLEAN.value,
        before="1.00",
        after="2.00",
    )
    base.update(overrides)
    return Resolution(**base)


def _document(*, claims=None, resolutions=None, **overrides) -> LossRunDocument:
    base = dict(
        source_filename="synthetic-loss-run.pdf",
        file_sha256=SHA,
        carrier="Synthetic Mutual",
        named_insured="Synthetic Holdings LLC",
        policy_number="POL-0001",
        valuation_date=datetime.date(2024, 1, 31),
        page_count=1,
        claims=claims if claims is not None else [_claim()],
        review_log=ReviewLog(entries=resolutions or []),
    )
    base.update(overrides)
    return LossRunDocument(**base)


def _doc_with_payload(where: str, payload: str) -> LossRunDocument:
    """Place one injection payload on one named workbook surface."""
    if where in ("claim_number", "claimant_name", "loss_description", "cause_of_loss"):
        return _document(claims=[_claim(**{where: payload})])
    if where in ("carrier", "named_insured", "policy_number", "source_filename"):
        return _document(**{where: payload})
    if where == "resolution.note":
        return _document(resolutions=[_resolution(note=payload)])
    if where == "resolution.before":
        return _document(resolutions=[_resolution(before=payload)])
    if where == "resolution.after":
        return _document(resolutions=[_resolution(after=payload)])
    raise ValueError(where)


FORMULA_SURFACES = [
    "claim_number",
    "claimant_name",
    "loss_description",
    "cause_of_loss",
    "carrier",
    "named_insured",
    "policy_number",
    "source_filename",
    "resolution.note",
    "resolution.before",
    "resolution.after",
]


def _formula_xlsx(where: str, payload: str) -> bytes:
    """Workbook bytes carrying the payload, including the finding.message surface.

    "Full detail" is used so every text surface named in the task — including
    claimant and cause — is actually present; the default template omits them.
    """
    if where == "finding.message":
        from core.reconcile import ReconciliationResult

        result = ReconciliationResult(
            status=DocumentStatus.NEEDS_REVIEW, findings=[_finding(message=payload)]
        )
        return to_bytes(_document(), result, template="Full detail")
    return to_bytes(_doc_with_payload(where, payload), template="Full detail")


ALL_FORMULA_CASES = [
    (where, payload)
    for where in FORMULA_SURFACES + ["finding.message"]
    for payload in FORMULA_PAYLOADS
]


def _zip_members(xlsx_bytes: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(xlsx_bytes)) as zf:
        return {name: zf.read(name) for name in zf.namelist()}


def _load(xlsx_bytes: bytes) -> openpyxl.Workbook:
    return openpyxl.load_workbook(io.BytesIO(xlsx_bytes))


def _cells_containing(wb: openpyxl.Workbook, needle: str):
    # A model validator strips surrounding whitespace from some fields (claim
    # numbers), so compare on the stripped value: the payload is still the
    # whole cell, and a leading TAB/space that survived must still be quoted.
    wanted = needle.strip()
    return [
        cell
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.strip() == wanted
    ]


# ---------------------------------------------------------------------------
# 1. Formula injection (spreadsheet formula execution on open)
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "where,payload",
    ALL_FORMULA_CASES,
    ids=[f"{w}::{i}" for w, _ in ALL_FORMULA_CASES for i in [ALL_FORMULA_CASES.index((w, _)) % 9]],
)
def test_formula_injection_no_formula_nodes(where, payload):
    """No <f> formula element may appear anywhere in the package."""
    members = _zip_members(_formula_xlsx(where, payload))
    for name, data in members.items():
        assert b"<f>" not in data, (
            f"formula element in {name} (payload {payload!r} via {where})"
        )


@pytest.mark.parametrize(
    "where,payload",
    ALL_FORMULA_CASES,
    ids=[f"{w}::{p!r}" for w, p in ALL_FORMULA_CASES],
)
def test_formula_injection_cells_are_quoted_strings(where, payload):
    """Payload cells must be stored as quoted strings, never as values Excel
    can interpret as formulas."""
    wb = _load(_formula_xlsx(where, payload))
    cells = _cells_containing(wb, payload)
    assert cells, f"payload {payload!r} missing from workbook (surface {where})"
    for cell in cells:
        assert cell.data_type == "s", (
            f"payload cell {cell.parent.title}!{cell.coordinate} must be a string, "
            f"got data_type={cell.data_type!r} (surface {where})"
        )
        assert cell.quotePrefix, (
            f"payload cell {cell.parent.title}!{cell.coordinate} must carry "
            f"quotePrefix so Excel shows it as text (surface {where})"
        )


# ---------------------------------------------------------------------------
# 2. Numeric/date preservation vs numeric-looking text
# ---------------------------------------------------------------------------


def _detail_cell(wb, header_name):
    ws = wb["Claim Detail"]
    header = [str(c.value).lower() for c in ws[1]]
    key = header_name.lower()
    assert key in header, f"{header_name!r} column not found: {header}"
    return ws.cell(row=2, column=header.index(key) + 1)


def test_negative_decimal_money_stays_numeric():
    doc = _document(claims=[_claim(paid_total=Decimal("-1500.25"))])
    wb = _load(to_bytes(doc))
    cell = _detail_cell(wb, "Paid Total")
    assert cell.data_type == "n", (
        f"negative money must stay numeric, got data_type={cell.data_type!r}"
    )
    assert float(cell.value) == pytest.approx(-1500.25)
    assert not cell.quotePrefix


def test_date_of_loss_stays_a_date_cell():
    wb = _load(to_bytes(_document()))
    cell = _detail_cell(wb, "Date of Loss")
    assert cell.is_date, f"date of loss must remain a date cell, got {cell.value!r}"
    assert not cell.quotePrefix


@pytest.mark.parametrize("text", ["-0001", "+44 20 7946 0000", "0123"])
def test_numeric_looking_text_stays_text(text):
    doc = _document(claims=[_claim(claim_number=text)])
    wb = _load(to_bytes(doc))
    cell = _detail_cell(wb, "Claim Number")
    assert cell.data_type == "s", (
        f"numeric-looking text {text!r} must stay a string, got "
        f"data_type={cell.data_type!r} value={cell.value!r}"
    )
    assert cell.value == text, f"leading zeros/signs must be preserved, got {cell.value!r}"


# ---------------------------------------------------------------------------
# 3. Control characters are scrubbed and the scrub is disclosed
# ---------------------------------------------------------------------------


def test_control_characters_replaced_and_disclosed():
    dirty = "Water damage in unit\x07 seven"
    doc = _document(claims=[_claim(loss_description=dirty)])
    xlsx = to_bytes(doc)  # must not raise
    wb = _load(xlsx)
    joined = "\n".join(
        str(v)
        for row in wb["Source Info"].iter_rows(values_only=True)
        for v in row
        if v is not None
    )
    assert "\x07" not in joined, "control character leaked into Source Info"
    assert re.search(r"replaced|control character", joined, re.IGNORECASE), (
        "Source Info must disclose that control characters were replaced"
    )
    assert re.search(r"\b1\b", joined), "replacement count (1) must appear in Source Info"


# ---------------------------------------------------------------------------
# 4. Redaction sentinels must never survive redact=True
# ---------------------------------------------------------------------------

SENT_CLAIMANT = "SENTINEL-CLAIMANT-QX7Z9"
SENT_DESC = "SENTINEL-DESC-WM3K8"
SENT_ORIG_CLAIMANT = "SENTINEL-ORIG-CLAIMANT-BV4N2"
SENT_ORIG_DESC = "SENTINEL-ORIG-DESC-JT6R1"
SENT_BEFORE = "SENTINEL-BEFORE-HG5D3"
SENT_AFTER = "SENTINEL-AFTER-LP9F7"
SENT_NOTE = f"note quoting claimant {SENT_CLAIMANT}"

ALL_SENTINELS = [
    SENT_CLAIMANT,
    SENT_DESC,
    SENT_ORIG_CLAIMANT,
    SENT_ORIG_DESC,
    SENT_BEFORE,
    SENT_AFTER,
    SENT_NOTE,
]


def _sentinel_document() -> LossRunDocument:
    claim = _claim(
        claimant_name=SENT_CLAIMANT,
        loss_description=SENT_DESC,
        original_values={
            "claimant_name": SENT_ORIG_CLAIMANT,
            "loss_description": SENT_ORIG_DESC,
        },
        edited_fields=["claimant_name", "loss_description"],
    )
    res = _resolution(
        note=SENT_NOTE,
        field="claimant_name",
        before=SENT_BEFORE,
        after=SENT_AFTER,
    )
    return _document(claims=[claim], resolutions=[res])


def _assert_no_sentinels(xlsx: bytes, label: str):
    for sentinel in ALL_SENTINELS:
        needle = sentinel.encode("utf-8")
        for name, data in _zip_members(xlsx).items():
            assert needle not in data, (
                f"sentinel {sentinel!r} survived redaction in {name} ({label})"
            )


def test_redaction_removes_sentinels_document_export():
    doc = _sentinel_document()
    _assert_no_sentinels(to_bytes(doc, redact=True), "to_bytes(redact=True)")
    buf = io.BytesIO()
    build_workbook(doc, redact=True).save(buf)
    _assert_no_sentinels(buf.getvalue(), "build_workbook(redact=True)")


def test_redaction_removes_sentinels_account_export():
    acct = build_account("Synthetic Account", [_sentinel_document()])
    _assert_no_sentinels(account_to_bytes(acct, redact=True), "account_to_bytes(redact=True)")
    buf = io.BytesIO()
    build_account_workbook(acct, redact=True).save(buf)
    _assert_no_sentinels(buf.getvalue(), "build_account_workbook(redact=True)")


def test_redaction_toggle_off_preserves_sentinels():
    """Guard against over-redaction: with redact=False the claim and review
    surfaces must still carry the data."""
    members = _zip_members(to_bytes(_sentinel_document(), redact=False))
    found = {
        s
        for s in ALL_SENTINELS
        if any(s.encode("utf-8") in data for data in members.values())
    }
    assert SENT_CLAIMANT in found
    assert SENT_DESC in found
    assert SENT_NOTE in found


# ---------------------------------------------------------------------------
# 5. Source Info disclosure of redaction state
# ---------------------------------------------------------------------------


def _source_info_rows(wb) -> dict:
    return {
        str(k): str(v)
        for k, v, *_ in wb["Source Info"].iter_rows(values_only=True)
        if k is not None
    }


def test_source_info_states_redaction_yes():
    wb = _load(to_bytes(_sentinel_document(), redact=True))
    rows = _source_info_rows(wb)
    assert "Claimant data redacted" in rows, "Source Info lacks redaction disclosure"
    assert rows["Claimant data redacted"] == "yes"


def test_source_info_states_redaction_no_when_off():
    wb = _load(to_bytes(_sentinel_document(), redact=False))
    rows = _source_info_rows(wb)
    assert rows.get("Claimant data redacted") == "no"


def test_redacted_fields_constant_names_personal_data_fields():
    assert set(REDACTED_FIELDS) == {"claimant_name", "loss_description"}


def test_suggested_filename_carries_no_personal_data():
    doc = _sentinel_document()
    name = suggested_filename(doc)
    for sentinel in ALL_SENTINELS:
        assert sentinel not in name
