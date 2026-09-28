"""XLSX export safety: no formula injection, no crash on control characters.

A loss run is untrusted input. These tests hand the exporter the payloads a
hostile carrier (or a hostile *filename*) could print, and demand that every
one is stored as literal text, that no ``<f>`` formula element exists anywhere
in the package, and that a value only looks numeric stays text.
"""

from __future__ import annotations

import io
import zipfile
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from core.export import DEFAULT_TEMPLATE, build_workbook, to_bytes
from core.schema import (
    Claim,
    Finding,
    FindingCategory,
    FindingScope,
    LossRunDocument,
    ReconciliationResult,
    ReviewAction,
    Resolution,
    Severity,
)

#: The payloads from the task. Every one can start a formula in Excel or a
#: spreadsheet library that trusts the leading character.
PAYLOADS = (
    "=cmd|'/c calc'!A0",
    '=HYPERLINK("http://x","y")',
    "+1+1",
    "-2+3",
    "@SUM(1)",
    "\t=1",
    "\r=1",
    "\n=1",
    " =1",
)


def _members(payload: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def hostile_document() -> tuple[LossRunDocument, ReconciliationResult]:
    """A document whose every text field carries a formula payload."""
    claims = [
        Claim(
            claim_number=PAYLOADS[0],
            claimant_name=PAYLOADS[1],
            loss_description=PAYLOADS[2],
            cause_of_loss=PAYLOADS[3],
        )
    ]
    document = LossRunDocument(
        source_filename=PAYLOADS[4],
        file_sha256="0" * 64,
        carrier=PAYLOADS[5],
        named_insured=PAYLOADS[6],
        policy_number=PAYLOADS[7],
        claims=claims,
    )
    document.review_log.record(
        Resolution(
            key="k1",
            action=ReviewAction.CORRECTED,
            rule_id="R-01",
            severity="ERROR",
            message=PAYLOADS[8],
            claim_number=PAYLOADS[0],
            field="loss_description",
            before=PAYLOADS[2],
            after=PAYLOADS[3],
            note=PAYLOADS[1],
        )
    )
    finding = Finding(
        rule_id="R-01",
        severity=Severity.ERROR,
        message=PAYLOADS[8],
        scope=FindingScope.DOCUMENT,
        category=FindingCategory.FINANCIAL,
        subject="document",
    )
    return document, ReconciliationResult(
        status="NEEDS_REVIEW", findings=[finding]
    )


def test_no_formula_element_survives_anywhere_in_the_package():
    document, result = hostile_document()
    payload = to_bytes(document, result, template="Full detail")
    for name, data in _members(payload).items():
        assert b"<f>" not in data, f"formula element in {name}"
        assert b"</f>" not in data, f"formula element in {name}"


def test_every_formula_payload_is_stored_as_quoted_text():
    document, result = hostile_document()
    payload = to_bytes(document, result, template="Full detail", include_provenance=True)
    workbook = load_workbook(io.BytesIO(payload))

    seen: set[str] = set()
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value in PAYLOADS:
                    seen.add(cell.value)
                    assert cell.data_type == "s", (sheet.title, cell.coordinate, cell.value)
                    assert cell.quotePrefix is True, (sheet.title, cell.coordinate, cell.value)

    assert seen == set(PAYLOADS), f"payloads not exercised: {set(PAYLOADS) - seen}"


def test_quote_prefix_is_present_in_the_raw_styles_part():
    document, result = hostile_document()
    payload = to_bytes(document, result, template="Full detail")
    assert b'quotePrefix="1"' in _members(payload)["xl/styles.xml"]


def test_negative_money_stays_numeric_and_dates_stay_dates():
    document = LossRunDocument(
        source_filename="numbers.pdf",
        file_sha256="h",
        claims=[
            Claim(
                claim_number="C1",
                incurred_total=Decimal("-1234.56"),
                date_of_loss=date(2024, 1, 1),
            )
        ],
    )
    workbook = build_workbook(document, None, template="Claim numbers and incurred")
    sheet = workbook["Claim Detail"]

    values = {cell.coordinate: cell for row in sheet.iter_rows() for cell in row}
    money = next(c for c in values.values() if c.value == float(Decimal("-1234.56")))
    assert money.data_type == "n"
    assert "#,##0.00" in money.number_format

    when = next(c for c in values.values() if isinstance(c.value, date))
    assert when.number_format.startswith("yyyy")


def test_text_that_only_looks_numeric_stays_text():
    document = LossRunDocument(
        source_filename="lookalikes.pdf",
        file_sha256="h",
        claims=[
            Claim(claim_number="-0001"),
            Claim(claim_number="+44 20 7946 0000"),
            Claim(claim_number="0123"),
        ],
    )
    workbook = build_workbook(document, None, template="Claim numbers and incurred")
    found = {
        cell.value: cell
        for row in workbook["Claim Detail"].iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value in {"-0001", "+44 20 7946 0000", "0123"}
    }
    assert set(found) == {"-0001", "+44 20 7946 0000", "0123"}
    for cell in found.values():
        assert cell.data_type == "s"


def test_control_character_does_not_crash_and_is_counted():
    document = LossRunDocument(
        source_filename="controls.pdf",
        file_sha256="h",
        claims=[Claim(claim_number="C1", loss_description="bell\x07here")],
    )
    payload = to_bytes(document, None, template="Full detail")

    for name, data in _members(payload).items():
        assert b"\x07" not in data, f"control character survived in {name}"

    workbook = load_workbook(io.BytesIO(payload))
    source = workbook["Source Info"]
    facts = {
        row[0].value: row[1].value
        for row in source.iter_rows(min_col=1, max_col=2)
        if row[0].value is not None
    }
    assert "Control characters replaced" in facts
    assert int(facts["Control characters replaced"]) >= 1

    body = "\n".join(
        str(cell.value)
        for row in workbook["Claim Detail"].iter_rows()
        for cell in row
        if isinstance(cell.value, str)
    )
    assert "\ufffd" in body
    assert "bell\x07here" not in body


def test_workbook_properties_pass_through_the_policy():
    document = LossRunDocument(
        source_filename="=1+1.pdf", file_sha256="h", claims=[Claim(claim_number="C1")]
    )
    workbook = build_workbook(document, None)
    # The subject carries the source filename; it must be a plain string with no
    # formula inference, and the property writer must have run.
    assert workbook.properties.creator == "LossLift"
    assert workbook.properties.subject == "=1+1.pdf"
