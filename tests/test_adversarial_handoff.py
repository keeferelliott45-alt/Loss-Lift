"""Fresh adversarial cases independent of the previous review fixtures."""

from datetime import date
from decimal import Decimal

import pytest

from core.pipeline import apply_edits, edit_claims, run_pipeline, to_records
from core.reconcile import reconcile
from core.schema import Claim, LossRunDocument, ReviewLog, SourceMethod, finding_key


@pytest.fixture
def result(golden_dir):
    return run_pipeline(golden_dir / "arithmetic_error.pdf", use_vision=False)


def test_two_date_order_failures_on_one_field_have_distinct_identities():
    doc = LossRunDocument(
        source_filename="dates.pdf", file_sha256="dates",
        valuation_date=date(2024, 12, 31),
        claims=[Claim(claim_number="001", date_of_loss=date(2025, 2, 1),
                      date_reported=date(2025, 1, 1), incurred_total=Decimal("0"))],
    )
    findings = [f for f in reconcile(doc).findings if f.rule_id == "R-10"]
    assert len(findings) == 2
    assert len({finding_key(f) for f in findings}) == 2


def test_missing_and_overlapping_source_metadata_never_duplicates_row_ids():
    claims = [
        Claim(claim_number="001", source_page=1, source_row=None),
        Claim(claim_number="001", source_page=1, source_row=None),
        Claim(claim_number="001", source_page=2, source_row=4, source_lines=[3, 4]),
        Claim(claim_number="001", source_page=2, source_row=4, source_lines=[4, 5]),
    ]
    doc = LossRunDocument(source_filename="rows.pdf", file_sha256="rows", claims=claims)
    ids = [c.row_id for c in doc.claims]
    assert len(set(ids)) == 4
    rebuilt = LossRunDocument.model_validate(doc.model_dump())
    assert [c.row_id for c in rebuilt.claims] == ids


def test_every_unshown_claim_field_survives_a_grid_edit(result):
    claim = result.document.claims[0]
    claim.claimant_ref = "000045"
    claim.loss_state = "CA"
    claim.close_date = date(2024, 5, 1)
    claim.body_part = "hand"
    claim.medical_only_flag = True
    claim.deductible_basis = "net"
    claim.alae_treatment = "separate"
    before = claim.model_dump()
    records = to_records(result.document)
    records[0]["loss_description"] = "reviewer clarification"
    changed = apply_edits(result.document, records).claims[0]
    for field in ("claimant_ref", "loss_state", "close_date", "body_part",
                  "medical_only_flag", "deductible_basis", "alae_treatment"):
        assert getattr(changed, field) == before[field]


def test_chronology_reconstructs_edits_reverts_and_repeated_renames(result):
    claim = result.document.claims[0]
    row_id, number, carrier = claim.row_id, claim.claim_number, claim.incurred_total
    raw, bbox = dict(claim.raw_cells), claim.source_bbox
    steps = [("incurred_total", "1.00"), ("incurred_total", "2.00"),
             ("incurred_total", str(carrier)), ("claim_number", "0000007"),
             ("incurred_total", "3.00"), ("claim_number", number)]
    for field, value in steps:
        records = to_records(result.document)
        next(r for r in records if r["_id"] == row_id)[field] = value
        edit_claims(result, records)
    entries = [e for e in result.document.review_log.entries if e.row_id == row_id]
    assert len(entries) == len(steps)
    current = {"claim_number": number, "incurred_total": str(carrier)}
    for entry, (field, value) in zip(entries, steps):
        assert entry.field == field
        assert entry.before == current[field]
        assert entry.after == value
        assert entry.status_before and entry.status_after
        current[field] = value
    assert [e.at for e in entries] == sorted(e.at for e in entries)
    claim = next(c for c in result.document.claims if c.row_id == row_id)
    assert claim.original_values["incurred_total"] == str(carrier)
    assert claim.raw_cells == raw and claim.source_bbox == bbox


@pytest.mark.parametrize("stage", ["normalization", "apply_edits", "reconciliation", "audit"])
def test_grid_failure_at_each_stage_preserves_the_previous_valid_state(result, monkeypatch, stage):
    import core.pipeline as pipeline

    before = result.document.model_dump(mode="json")
    reconciliation = result.reconciliation
    records = to_records(result.document)
    records[0]["incurred_total"] = "123.00"

    def fail(*args, **kwargs):
        raise RuntimeError(stage)

    if stage == "audit":
        monkeypatch.setattr(ReviewLog, "record", fail)
    else:
        monkeypatch.setattr(pipeline, {
            "normalization": "_coerce", "apply_edits": "apply_edits",
            "reconciliation": "rerun_reconciliation",
        }[stage], fail)
    with pytest.raises(RuntimeError, match=stage):
        edit_claims(result, records)
    assert result.document.model_dump(mode="json") == before
    assert result.reconciliation is reconciliation


def test_manual_row_identity_survives_renames_and_deletion_is_refused(result):
    records = to_records(result.document)
    records.append({"claim_number": "000001", "incurred_total": "12.50"})
    edit_claims(result, records)
    manual = result.document.claims[-1]
    row_id = manual.row_id
    assert manual.source_method is SourceMethod.MANUAL
    for number, money in [("000002", "13.50"), ("000001", "14.50")]:
        records = to_records(result.document)
        record = next(r for r in records if r["_id"] == row_id)
        record.update(claim_number=number, incurred_total=money)
        edit_claims(result, records)
        assert result.document.claims[-1].row_id == row_id
    before = result.document.model_dump(mode="json")
    with pytest.raises(ValueError, match="cannot be deleted"):
        edit_claims(result, to_records(result.document)[:-1])
    assert result.document.model_dump(mode="json") == before


def test_rereading_same_pdf_is_a_distinct_unreviewed_document(result, golden_dir):
    records = to_records(result.document)
    records[0]["incurred_total"] = "1.00"
    edit_claims(result, records)
    fresh = run_pipeline(golden_dir / "arithmetic_error.pdf", use_vision=False)
    assert fresh.document.file_sha256 == result.document.file_sha256
    assert fresh.document.document_id != result.document.document_id
    assert fresh.document.review_log.entries == []
    assert result.document.review_log.entries


def test_explicit_null_reason_change_is_audited_but_untouched_null_is_not(result):
    from core.schema import NullReason

    claim = result.document.claims[0]
    claim.incurred_total = None
    claim.field_issues["incurred_total"] = NullReason.UNPARSEABLE
    claim.raw_cells["incurred_total"] = "illegible"
    untouched = apply_edits(result.document, to_records(result.document))
    assert untouched.claims[0].field_issues["incurred_total"] is NullReason.UNPARSEABLE
    assert not untouched.review_log.entries
    records = to_records(result.document)
    records[0]["incurred_total"] = "N/A"
    edit_claims(result, records)
    changed = result.document.claims[0]
    assert changed.field_issues["incurred_total"] is NullReason.NOT_APPLICABLE
    assert changed.raw_cells["incurred_total"] == "illegible"
    entry = result.document.review_log.entries[-1]
    assert entry.before == "null (UNPARSEABLE)"
    assert entry.after == "null (NOT_APPLICABLE)"


def test_export_failure_leaves_the_document_and_history_unchanged(result, monkeypatch):
    from core.export import to_bytes
    from openpyxl.workbook.workbook import Workbook

    before = result.document.model_dump(mode="json")

    def fail(*args, **kwargs):
        raise RuntimeError("export failed")

    monkeypatch.setattr(Workbook, "save", fail)
    with pytest.raises(RuntimeError, match="export failed"):
        to_bytes(result.document, result.reconciliation)
    assert result.document.model_dump(mode="json") == before


def test_unchanged_canonical_decimal_text_is_not_reinterpreted_as_eu_thousands(result):
    result.document.locale_hint = "eu"
    result.document.claims[0].incurred_total = Decimal("1.234")
    updated = apply_edits(result.document, to_records(result.document))
    assert updated.claims[0].incurred_total == Decimal("1.234")
    assert not updated.review_log.entries


def test_unchanged_grid_does_not_publish_a_new_document(result):
    import pandas as pd

    original = result.document
    reconciliation = result.reconciliation
    records = pd.DataFrame(to_records(original)).to_dict("records")
    edit_claims(result, records)
    assert result.document is original
    assert result.reconciliation is reconciliation


def test_duplicate_number_group_identity_changes_when_physical_members_change():
    def document(rows):
        return LossRunDocument(source_filename="x.pdf", file_sha256="x", claims=rows)

    first = document([Claim(claim_number="DUP", row_id="p1r1"),
                      Claim(claim_number="DUP", row_id="p1r2")])
    second = document([Claim(claim_number="DUP", row_id="p1r1"),
                       Claim(claim_number="DUP", row_id="p1r3")])
    one = next(f for f in reconcile(first).findings if f.rule_id == "R-11")
    two = next(f for f in reconcile(second).findings if f.rule_id == "R-11")
    assert one.actual == two.actual
    assert finding_key(one) != finding_key(two)
