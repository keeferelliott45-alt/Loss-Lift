"""The submission workbook: formula-safe, fully redactable, traceable.

Synthetic emails and PDFs only.
"""

from __future__ import annotations

import io
import zipfile

import openpyxl
import pytest

from core.eml_intake import read_submission
from core.ingest import discard
from core.submission import summarise_submission
from core.submission_export import (
    SHEETS,
    build_submission_workbook,
    submission_filename,
    submission_to_bytes,
)
from tests.submission_fixtures import BODY_SECRET, claim_rows, eml, loss_run_pdf, run

FORMULA = "=cmd|' calc'!A0"
SUBJECT_FORMULA = "@SUM(1+1) losses"
SENDER_SENTINEL = "Sentinel Sender <sender-sentinel@example.test>"
SUBJECT_SENTINEL = "SUBJECT-SENTINEL renewal losses"
FILE_SENTINEL = "FILENAME-SENTINEL-Harbor.pdf"
INSURED_SENTINEL = "Insuredsentinel Holdings LLC"
POLICY_SENTINEL = "POLSENT-778899"


@pytest.fixture
def submission(tmp_path):
    pdf = loss_run_pdf(insured=INSURED_SENTINEL, policy=POLICY_SENTINEL,
                       rows=claim_rows(3, big=60000))
    raw = eml(
        [(FILE_SENTINEL, "application/pdf", pdf),
         ("junk.zip", "application/zip", b"PK\x03\x04" + b"\0" * 20),
         (f"{FORMULA}.pdf", "application/pdf", b"not a pdf")],
        sender=SENDER_SENTINEL, subject=SUBJECT_SENTINEL,
    )
    sub, done = read_submission(raw, run(tmp_path / "profiles"))
    results = {doc_id: result for doc_id, (_s, result) in done.items()}
    yield sub, summarise_submission(sub, results), results
    for staged, _r in done.values():
        discard(staged)


def _members(payload: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def test_every_sheet_is_present_and_nothing_is_hidden_or_linked(submission):
    sub, summary, results = submission
    payload = submission_to_bytes(sub, summary, results)
    workbook = openpyxl.load_workbook(io.BytesIO(payload))
    assert tuple(workbook.sheetnames) == SHEETS
    assert all(sheet.sheet_state == "visible" for sheet in workbook.worksheets)
    members = _members(payload)
    for name, data in members.items():
        assert b"<f>" not in data, name
        assert b"<hyperlink" not in data, name
        assert not name.startswith("xl/externalLinks"), name
    assert not any(b"TargetMode=\"External\"" in data for data in members.values())


def test_the_email_body_is_never_exported(submission):
    sub, summary, results = submission
    for redact in (False, True):
        payload = submission_to_bytes(sub, summary, results, redact=redact)
        assert all(BODY_SECRET.encode() not in data for data in _members(payload).values())


def test_formula_payloads_are_stored_as_quoted_text(tmp_path):
    raw = eml([(f"{FORMULA}.pdf", "application/pdf", loss_run_pdf())],
              subject=SUBJECT_FORMULA, sender="+1+1 <plus@example.test>")
    sub, done = read_submission(raw, run(tmp_path / "profiles"))
    try:
        results = {d: r for d, (_s, r) in done.items()}
        payload = submission_to_bytes(sub, summarise_submission(sub, results), results)
        workbook = openpyxl.load_workbook(io.BytesIO(payload))
        cells = [c for s in workbook.worksheets for r in s.iter_rows() for c in r
                 if isinstance(c.value, str) and c.value.startswith((FORMULA, "@SUM", "+1+1"))]
        found = {c.value[:4] for c in cells}
        assert {FORMULA[:4], "@SUM", "+1+1"} <= found
        assert all(c.data_type == "s" and c.quotePrefix for c in cells)
        assert all(b"<f>" not in data for data in _members(payload).values())
    finally:
        for staged, _r in done.values():
            discard(staged)


def test_unredacted_export_carries_the_inventory_and_provenance(submission):
    sub, summary, results = submission
    workbook = build_submission_workbook(sub, summary, results)
    attachments = list(workbook["Attachments"].iter_rows(min_row=2, values_only=True))
    assert [row[1] for row in attachments] == [a.attachment_id for a in sub.attachments]
    assert [row[5] for row in attachments] == [a.sha256 for a in sub.attachments]
    assert [row[6] for row in attachments] == ["Accepted", "Rejected", "Rejected"]
    large = list(workbook["Large Claims"].iter_rows(min_row=2, values_only=True))
    (claim,) = summary.accounts[0].large_claims
    assert large[0][1] == claim.claim_number
    assert large[0][5:10] == (claim.provenance.attachment_id, claim.provenance.document_id,
                              claim.provenance.run_id, claim.provenance.page,
                              claim.provenance.row)
    values = {r[0]: r[1] for r in workbook["Submission"].iter_rows(min_row=2, values_only=True)}
    assert values["Submission ID"] == sub.submission_id
    assert values["Sender"] == SENDER_SENTINEL
    assert values["Status"].startswith("Incomplete")


SENTINELS = [SENDER_SENTINEL, "sender-sentinel", SUBJECT_SENTINEL, "SUBJECT-SENTINEL",
             FILE_SENTINEL, "FILENAME-SENTINEL", INSURED_SENTINEL, POLICY_SENTINEL]


def test_redaction_removes_identifying_values_from_every_part(submission):
    sub, summary, results = submission
    payload = submission_to_bytes(sub, summary, results, redact=True)
    members = _members(payload)
    for name, data in members.items():
        text = data.decode("utf-8", errors="ignore")
        for sentinel in SENTINELS:
            assert sentinel.lower() not in text.lower(), (name, sentinel)
    # Audit identifiers and hashes survive redaction.
    joined = b"".join(members.values()).decode("utf-8", errors="ignore")
    assert sub.submission_id in joined
    assert all(a.attachment_id in joined for a in sub.attachments)
    assert all(a.sha256 in joined for a in sub.attachments if a.sha256)


def test_redaction_off_keeps_the_values(submission):
    sub, summary, results = submission
    joined = b"".join(_members(submission_to_bytes(sub, summary, results)).values())
    text = joined.decode("utf-8", errors="ignore")
    assert SENDER_SENTINEL.split("<")[0].strip() in text
    assert INSURED_SENTINEL in text


def test_the_workbook_name_carries_nothing_from_the_email(submission):
    sub, _summary, _results = submission
    name = submission_filename(sub)
    assert name == f"losslift-{sub.submission_id}.xlsx"
    assert "SENTINEL" not in name.upper()


def test_a_short_subject_or_file_name_never_scrubs_labels_or_headers(tmp_path):
    raw = eml([("Claim.pdf", "application/pdf", loss_run_pdf())], subject="Claim",
              sender="Re <re@example.test>")
    sub, done = read_submission(raw, run(tmp_path / "profiles"))
    try:
        results = {d: r for d, (_s, r) in done.items()}
        workbook = build_submission_workbook(sub, summarise_submission(sub, results),
                                             results, redact=True)
        headers = [c.value for c in workbook["Claims"][1]]
        assert "Claim number" in headers
        labels = [r[0] for r in workbook["Submission"].iter_rows(min_row=2, values_only=True)]
        assert "Submission ID" in labels and "Every attachment accounted for" in labels
        values = {r[0]: r[1] for r in workbook["Submission"].iter_rows(min_row=2,
                                                                         values_only=True)}
        assert values["Subject"] == "[redacted]" and values["Sender"] == "[redacted]"
    finally:
        for staged, _r in done.values():
            discard(staged)


def test_an_email_date_without_a_zone_is_read_as_utc():
    from datetime import datetime

    from core.submission_export import _utc

    assert _utc(datetime(2026, 9, 29, 10, 15)) == "2026-09-29 10:15:00 UTC"


def test_an_insured_printed_with_extra_spacing_is_still_redacted(submission):
    """The account rollup collapses spacing; a blocker quoting it must still scrub."""
    from core.submission_export import _identifying

    sub, _summary, results = submission
    (result,) = results.values()
    result.document.named_insured = "Doublespace   Sentinel  Holdings LLC"
    scrubbed = _identifying(sub, results).scrub(
        "Doublespace Sentinel Holdings LLC: 2 claim(s) listed in an earlier valuation")
    assert "sentinel" not in scrubbed.lower() and scrubbed.startswith("[redacted]")


def test_every_counted_claim_is_exported_with_where_it_was_read(submission):
    sub, summary, results = submission
    workbook = build_submission_workbook(sub, summary, results)
    rows = list(workbook["Claims"].iter_rows(min_row=2, values_only=True))
    lines = [line for account in summary.accounts for line in account.claim_lines]
    assert len(rows) == len(lines) == summary.claims
    for row, line in zip(rows, lines):
        assert row[1] == line.claim_number
        assert row[10:15] == (line.provenance.attachment_id, line.provenance.document_id,
                              line.provenance.run_id, line.provenance.page,
                              line.provenance.row)
    headers = [c.value for c in workbook["Accounts"][1]]
    read_from = workbook["Accounts"].cell(row=2, column=headers.index(
        "Read from (attachment ID / run)") + 1).value
    assert read_from == sub.attachments[0].attachment_id
