"""A logical run is judged by what its own pages print.

Two carriers' reports bound in one PDF have two carriers, two policies, two
terms and two valuation dates. The rules run over each run with that run's
facts (``core.reconcile.run_view``), and the workbook writes each claim's row
with its own run's facts -- never the first report's letterhead standing in
for every page. Synthetic documents only (spec section 9).
"""

from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from core.export import to_bytes
from core.review import canonical_run_status
from core.schema import DocumentStatus
from tests.test_packet_adversarial import _read, _total
from tests.test_packet_claim_series import LARGE_CARRIER, SMALL_CARRIER, SMALL_HEADERS

GL_TERM = ("Policy Number: GL-100", "Policy Period: 01/01/2022 to 01/01/2023")
AUTO_TERM = ("Policy Number: AU-200", "Policy Period: 01/01/2020 to 01/01/2021")

RUN_ONE = [
    ("71004410", "03/10/2022", "CLOSED", "1,000.00", "0.00", "1,000.00"),
    ("71004411", "07/22/2022", "CLOSED", "2,500.00", "0.00", "2,500.00"),
]
RUN_TWO = [
    ("CR-50001", "03/15/2020", "CLOSED", "400.00", "0.00", "400.00"),
    ("CR-50002", "09/01/2020", "CLOSED", "600.00", "0.00", "600.00"),
    ("CR-50003", "05/05/2021", "CLOSED", "700.00", "0.00", "700.00"),   # after its own term
]


def _packet(tmp_path, *, second_valuation=True, second_carrier=True):
    second_top = (
        "Page 1 of 1",
        *((SMALL_CARRIER,) if second_carrier else ()),
        "LOSS RUN REPORT",
        *AUTO_TERM,
        *(("Valuation Date: 06/30/2021",) if second_valuation else ()),
    )
    return _read(tmp_path, [
        {"top": ("Page 1 of 1", LARGE_CARRIER, "LOSS RUN REPORT", *GL_TERM,
                 "Valuation Date: 12/31/2022"),
         "rows": RUN_ONE, "total": _total(RUN_ONE)},
        {"top": second_top, "headers": SMALL_HEADERS, "rows": RUN_TWO,
         "total": _total(RUN_TWO)},
    ])


def _claim_rows(result) -> dict[str, dict]:
    sheet = load_workbook(BytesIO(to_bytes(result.document, result.reconciliation)))["Claim Detail"]
    header = [cell.value for cell in sheet[1]]
    rows = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        record = dict(zip(header, row))
        rows[record["Claim number"]] = record
    return rows


def test_each_run_keeps_its_own_carrier_policy_term_and_valuation(tmp_path):
    result = _packet(tmp_path)
    one, two = result.document.runs
    assert (one.carrier, one.policy_number) == (LARGE_CARRIER, "GL-100")
    assert (two.carrier, two.policy_number) == (SMALL_CARRIER, "AU-200")
    assert (one.policy_period_start, one.policy_period_end) == (date(2022, 1, 1), date(2023, 1, 1))
    assert (two.policy_period_start, two.policy_period_end) == (date(2020, 1, 1), date(2021, 1, 1))
    assert (one.valuation_date, two.valuation_date) == (date(2022, 12, 31), date(2021, 6, 30))


def test_loss_dates_are_judged_against_their_own_runs_term(tmp_path):
    """Run 2's 2020 claims sit inside run 2's term. Judged against run 1's
    2022 term every one of them was out of period; only the one after run
    2's own term is."""
    result = _packet(tmp_path)
    outside = [f for f in result.reconciliation.findings if f.rule_id == "R-09"]
    assert [(f.claim_number, f.run_id) for f in outside] == [("CR-50003", "run-2")]
    assert outside[0].expected == "2020-01-01 to 2021-01-01"


def test_each_claim_row_carries_its_own_runs_facts(tmp_path):
    rows = _claim_rows(_packet(tmp_path))
    first, second = rows["71004410"], rows["CR-50001"]
    assert (first["Run ID"], first["Carrier"], first["Policy number"]) \
        == ("run-1", LARGE_CARRIER, "GL-100")
    assert (second["Run ID"], second["Carrier"], second["Policy number"]) \
        == ("run-2", SMALL_CARRIER, "AU-200")
    assert second["Policy term start"].date() == date(2020, 1, 1)
    assert second["Valuation date"].date() == date(2021, 6, 30)


def test_a_run_that_prints_no_valuation_date_does_not_borrow_one(tmp_path):
    """R-06 is raised for the run that prints none; its rows say nothing
    rather than run 1's date; run 1 is untouched."""
    result = _packet(tmp_path, second_valuation=False)
    reconciliation = result.reconciliation
    missing = [f for f in reconciliation.findings if f.rule_id == "R-06"]
    assert [f.run_id for f in missing] == ["run-2"]
    assert canonical_run_status(reconciliation, "run-1") is DocumentStatus.CLEAN
    assert canonical_run_status(reconciliation, "run-2") is DocumentStatus.NEEDS_REVIEW
    assert _claim_rows(result)["CR-50001"]["Valuation date"] is None


def test_a_run_that_prints_no_carrier_is_not_given_another_runs(tmp_path):
    """Whatever the run's own page yields for a carrier, it is never run 1's.
    (``detect_carrier`` can take a column-label line for a name when no
    letterhead prints one -- a known, document-level limitation.)"""
    result = _packet(tmp_path, second_carrier=False)
    assert result.document.runs[1].carrier != LARGE_CARRIER
    assert _claim_rows(result)["CR-50001"]["Carrier"] != LARGE_CARRIER
