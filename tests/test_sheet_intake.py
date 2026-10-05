"""Stage F — CSV/XLSX intake.

Every visible non-empty worksheet is a separate source and its own logical
run; formulas are never executed; nothing ambiguous becomes zero. Synthetic
openpyxl fixtures and in-memory CSV only.
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from decimal import Decimal

import pytest
from openpyxl import Workbook

from core.ingest import IngestError, ingest
from core.pipeline import run_pipeline
from core.review import canonical_status
from core.schema import DocumentStatus, FindingScope, NullReason, SourceMethod


HEADERS = [
    "Claim Number",
    "Date of Loss",
    "Claim Status",
    "Paid",
    "Reserve",
    "Incurred",
]


def _row(number, dol="05/23/2023", status="Open", paid="1,500.00",
         reserve="500.00", incurred="2,000.00"):
    return [number, dol, status, paid, reserve, incurred]


def _workbook(sheets: dict[str, list[list[str]]], hidden: dict[str, list[list[str]]] | None = None) -> bytes:
    buffer = io.BytesIO()
    book = Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        sheet = book.create_sheet(name[:31])
        for row in rows:
            sheet.append(row)
    for name, rows in (hidden or {}).items():
        sheet = book.create_sheet(name[:31])
        sheet.sheet_state = "hidden"
        for row in rows:
            sheet.append(row)
    book.save(buffer)
    return buffer.getvalue()


def _cached_formula(xlsx: bytes, sheet_name: str, cell: str, value: str) -> bytes:
    """Store a cached result under a formula, as Excel itself would.

    openpyxl cannot write a cached value, so the workbook XML is patched the
    same way a saved Excel workbook stores it. Nothing is evaluated here:
    the ``<v>`` is literal data, which is exactly what ``data_only`` reads.
    """
    source = zipfile.ZipFile(io.BytesIO(xlsx))
    workbook_xml = source.read("xl/workbook.xml").decode("utf-8")
    match = re.search(
        r'<sheet[^>]*name="' + re.escape(sheet_name) + r'"[^>]*r:id="(rId\d+)"',
        workbook_xml,
    )
    assert match, f"sheet {sheet_name!r} not found"
    rels = source.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    relationships = {
        id_match.group(1): target_match.group(1)
        for id_match, target_match in (
            (re.search(r'Id="(rId\d+)"', tag), re.search(r'Target="([^"]+)"', tag))
            for tag in re.findall(r"<Relationship\b[^>]*/>", rels)
        )
        if id_match and target_match
    }
    target = relationships[match.group(1)]
    if target.startswith("/"):
        target = target.lstrip("/")
    elif not target.startswith("xl/"):
        target = "xl/" + target
    xml = source.read(target).decode("utf-8")
    cell_xml = re.search(r'<c r="' + cell + r'"[^>]*>.*?</c>', xml)
    assert cell_xml and "<f>" in cell_xml.group(0), f"formula cell {cell} not found"
    patched = xml.replace(
        cell_xml.group(0),
        cell_xml.group(0).replace("<v></v>", f"<v>{value}</v>"),
    )
    assert patched != xml, "cached value was not inserted"
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target_zip:
        for item in source.infolist():
            payload = source.read(item.filename)
            if item.filename == target:
                payload = patched.encode("utf-8")
            target_zip.writestr(item, payload)
    return out.getvalue()


def _csv(rows: list[list[str]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8")


def _staged(tmp_path, data: bytes, name: str):
    return ingest(data, name, workdir=tmp_path)


def _result(tmp_path, data: bytes, name: str):
    return run_pipeline(_staged(tmp_path, data, name), use_vision=False)


# ---------------------------------------------------------------------------
# Ingest: kind sniffing and staging
# ---------------------------------------------------------------------------


def test_ingest_accepts_xlsx_and_csv_and_stages_the_right_extension(tmp_path):
    xlsx = _staged(tmp_path, _workbook({"Run": [HEADERS, _row("C-1")]}), "run.xlsx")
    assert xlsx.kind == "xlsx"
    assert xlsx.path.suffix == ".xlsx"
    csv_file = _staged(tmp_path, _csv([HEADERS, _row("C-1")]), "run.csv")
    assert csv_file.kind == "csv"
    assert csv_file.path.suffix == ".csv"


def test_ingest_rejects_extension_magic_mismatch(tmp_path):
    with pytest.raises(IngestError):
        _staged(tmp_path, _workbook({"Run": [HEADERS]}), "run.pdf")
    with pytest.raises(IngestError):
        _staged(tmp_path, b"%PDF-1.7 not a workbook", "run.xlsx")
    with pytest.raises(IngestError):
        _staged(tmp_path, _csv([HEADERS]), "run.xlsx")


def test_ingest_rejects_legacy_and_unknown_formats_fail_closed(tmp_path):
    # Legacy .xls (BIFF inside an OLE container) and a zip that is no workbook.
    ole = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 32
    with pytest.raises(IngestError):
        _staged(tmp_path, ole, "legacy.xls")
    with pytest.raises(IngestError):
        _staged(tmp_path, ole, "legacy.xlsx")
    zip_only = io.BytesIO()
    with zipfile.ZipFile(zip_only, "w") as archive:
        archive.writestr("readme.txt", "not a workbook")
    with pytest.raises(IngestError):
        _staged(tmp_path, zip_only.getvalue(), "junk.xlsx")


# ---------------------------------------------------------------------------
# Per-sheet sources and logical runs
# ---------------------------------------------------------------------------


def test_two_visible_sheets_are_two_sources_and_two_runs(tmp_path):
    result = _result(tmp_path, _workbook({
        "West": [HEADERS, _row("C-1001"), _row("C-1002")],
        "East": [HEADERS, _row("C-2001")],
    }), "packet.xlsx")
    document = result.document
    assert [run.run_id for run in document.runs] == ["run-1", "run-2"]
    assert document.runs[0].pages == [1]
    assert document.runs[1].pages == [2]
    assert document.is_packet
    numbers_by_run = {
        run.run_id: sorted(claim.claim_number for claim in document.run_claims(run))
        for run in document.runs
    }
    assert numbers_by_run == {"run-1": ["C-1001", "C-1002"], "run-2": ["C-2001"]}
    east_claim = document.run_claims(document.runs[1])[0]
    assert east_claim.source_page == 2
    assert east_claim.source_method is SourceMethod.SHEET


def test_single_sheet_workbook_is_one_report_with_no_runs(tmp_path):
    result = _result(tmp_path, _workbook({"Run": [HEADERS, _row("C-1001")]}), "single.xlsx")
    document = result.document
    assert document.runs == []
    assert [claim.claim_number for claim in document.claims] == ["C-1001"]
    assert document.claims[0].source_page == 1


def test_csv_upload_is_one_source(tmp_path):
    result = _result(tmp_path, _csv([HEADERS, _row("C-3001")]), "run.csv")
    document = result.document
    assert document.runs == []
    assert [claim.claim_number for claim in document.claims] == ["C-3001"]
    assert document.claims[0].source_method is SourceMethod.SHEET


def test_empty_visible_sheet_is_recorded_never_a_zero_run(tmp_path):
    result = _result(tmp_path, _workbook({
        "West": [HEADERS, _row("C-1001")],
        "Spare": [],
    }), "with-empty.xlsx")
    document = result.document
    assert document.ignored_sheets.get("Spare") == "empty"
    assert all(claim.source_page == 1 for claim in document.claims)
    # One visible non-empty sheet: a single loss run, and the empty sheet is
    # a recorded fact on the document, not a clean empty run.
    assert document.runs == []
    assert any("Spare" in warning for warning in result.warnings)


def test_hidden_sheet_is_ignored_but_recorded(tmp_path):
    result = _result(tmp_path, _workbook(
        {"West": [HEADERS, _row("C-1001")]},
        hidden={"Archive 2023": [HEADERS, _row("OLD-1")]},
    ), "with-hidden.xlsx")
    document = result.document
    assert document.ignored_sheets.get("Archive 2023") == "hidden"
    assert [claim.claim_number for claim in document.claims] == ["C-1001"]


# ---------------------------------------------------------------------------
# Formulas: cached value used, uncached refused
# ---------------------------------------------------------------------------


def test_formula_cached_result_is_used_never_evaluated(tmp_path):
    book = _workbook({"Run": [HEADERS, ["C-1001", "05/23/2023", "Open",
                                        "=SUM(1,2)", "500.00", "2503.00"]]})
    data = _cached_formula(book, "Run", "D2", "2003.00")
    result = _result(tmp_path, data, "cached.xlsx")
    claim = result.document.claims[0]
    assert claim.paid_total == Decimal("2003.00")
    assert claim.field_issues.get("paid_total") is None
    assert claim.raw_cells["paid_total"] == "2003"


def test_formula_without_cached_result_is_null_finding_needs_review(tmp_path):
    result = _result(tmp_path, _workbook({
        "Run": [HEADERS, ["C-1001", "05/23/2023", "Open",
                          "=SUM(1,2)", "500.00", "2503.00"]],
    }), "uncached.xlsx")
    claim = result.document.claims[0]
    assert claim.paid_total is None
    assert claim.field_issues.get("paid_total") is NullReason.UNPARSEABLE
    assert claim.raw_cells.get("paid_total") == "=SUM(1,2)"
    assert any(f.rule_id == "R-15" and f.claim_number == "C-1001"
               for f in result.reconciliation.findings)
    assert canonical_status(result.reconciliation) is DocumentStatus.NEEDS_REVIEW


# ---------------------------------------------------------------------------
# Deterministic parsing of numeric-looking text
# ---------------------------------------------------------------------------


def test_numeric_looking_text_is_parsed_deterministically(tmp_path):
    result = _result(tmp_path, _workbook({"Run": [
        HEADERS,
        _row("C-1001", paid="4,660", reserve="0.00", incurred="4,660"),
    ]}), "numbers.xlsx")
    claim = result.document.claims[0]
    assert claim.paid_total == Decimal("4660")
    assert claim.date_of_loss is not None and claim.date_of_loss.isoformat() == "2023-05-23"
    assert claim.claim_status.value == "OPEN"


# ---------------------------------------------------------------------------
# Headers: duplicates, blanks and unrecognised labels are flagged, not guessed
# ---------------------------------------------------------------------------


def test_duplicate_and_blank_headers_are_flagged_not_guessed(tmp_path):
    result = _result(tmp_path, _workbook({"Run": [
        ["Claim Number", "Paid", "Paid", "Incurred"],
        ["C-1001", "1,500.00", "2,000.00", "2,000.00"],
    ]}), "duplicate.xlsx")
    assert any("duplicate" in warning.lower() for warning in result.warnings)
    assert any(f.rule_id == "R-21" for f in result.reconciliation.findings)
    assert canonical_status(result.reconciliation) is DocumentStatus.NEEDS_REVIEW


def test_sheet_with_no_header_row_is_flagged_and_reads_no_claims(tmp_path):
    result = _result(tmp_path, _workbook({"Run": [
        _row("C-1001"),
    ]}), "noheader.xlsx")
    assert result.document.claims == []
    assert any("header" in warning.lower() for warning in result.warnings)
    assert any(f.rule_id == "R-20" for f in result.reconciliation.findings)


# ---------------------------------------------------------------------------
# Accountability: every claim-like row is read or recorded
# ---------------------------------------------------------------------------


def test_row_with_money_and_no_claim_number_is_recorded_not_dropped(tmp_path):
    result = _result(tmp_path, _workbook({"Run": [
        HEADERS,
        _row("C-1001"),
        ["", "06/01/2024", "Closed", "900.00", "", "900.00"],
    ]}), "orphan.xlsx")
    document = result.document
    assert len(document.unplaced_rows) == 1
    assert document.unplaced_rows[0].amounts
    assert any(f.rule_id == "R-23" for f in result.reconciliation.findings)


# ---------------------------------------------------------------------------
# Provenance and deterministic replay
# ---------------------------------------------------------------------------


def test_provenance_points_at_the_sheet_cell(tmp_path):
    result = _result(tmp_path, _workbook({
        "West": [["Title line"],
                 ["Claim Number", "Paid", "Incurred"],
                 ["C-1001", "500.00", "500.00"]],
    }), "prov.xlsx")
    claim = result.document.claims[0]
    assert claim.source_page == 1
    assert claim.source_row == 2  # 0-based: title, header, then this row
    assert claim.source_method is SourceMethod.SHEET


def test_replay_reproduces_the_same_claims(tmp_path):
    data = _workbook({
        "West": [HEADERS, _row("C-1001"), _row("C-1002")],
        "East": [HEADERS, _row("C-2001")],
    })

    def read():
        result = _result(tmp_path, data, "replay.xlsx")
        return sorted(
            (claim.claim_number, claim.source_page, claim.source_row,
             claim.paid_total, claim.incurred_total)
            for claim in result.document.claims
        )

    assert read() == read()
