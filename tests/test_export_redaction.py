"""Redaction covers the whole workbook, not just the claim columns.

The redaction toggle drops claimant names and loss descriptions. That was only
ever true of the Claim Detail sheet: review decisions, reviewer notes and
finding text were written whatever the setting. These tests put a unique
sentinel in every place a sensitive value can reach a workbook — the columns,
``Claim.original_values``, a correction's before/after, a reviewer's note, a
finding message — and grep **every** ZIP member of every export for it.
"""

from __future__ import annotations

import io
import zipfile

from core.account import build_account
from core.export import (
    account_to_bytes,
    build_batch_zip,
    build_workbook,
    to_bytes,
)
from core.schema import (
    Claim,
    Finding,
    FindingCategory,
    FindingScope,
    LossRunDocument,
    ReconciliationResult,
    ReviewAction,
    Resolution,
    Severity,
)

CLAIMANT = "Zyx Sentinelclaimant"
DESC = "Sentinel loss description 4471"
ORIG_NAME = "Sentinel original name 991"
ORIG_DESC = "Sentinel original desc 771"
BEFORE = "Sentinel before 552"
AFTER = "Sentinel after 553"
NOTE = f"reviewer note quoting {CLAIMANT} and {DESC}"
MESSAGE = f"finding message mentioning {CLAIMANT}"

SENTINELS = (CLAIMANT, DESC, ORIG_NAME, ORIG_DESC, BEFORE, AFTER, NOTE, MESSAGE)


def _members(payload: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def sentinel_document() -> tuple[LossRunDocument, ReconciliationResult]:
    claim = Claim(
        claim_number="S-1",
        claimant_name=CLAIMANT,
        loss_description=DESC,
        # A field that redaction does not drop, but which quotes a sensitive
        # value, to exercise the substring scrub in a surviving cell.
        cause_of_loss=f"cause about {ORIG_NAME}",
        original_values={
            "claimant_name": ORIG_NAME,
            "loss_description": ORIG_DESC,
        },
    )
    document = LossRunDocument(
        source_filename=f"{CLAIMANT}.pdf",
        file_sha256="0" * 64,
        named_insured="Sentinel Insured LLC",
        # Document-level field that redaction keeps, quoting an original value.
        policy_number=ORIG_DESC,
        carrier="Sentinel Carrier",
        claims=[claim],
    )
    document.review_log.record(
        Resolution(
            key="k1",
            action=ReviewAction.CORRECTED,
            rule_id="R-01",
            severity="ERROR",
            message=MESSAGE,
            claim_number="S-1",
            field="claimant_name",
            before=BEFORE,
            after=AFTER,
            note=NOTE,
        )
    )
    finding = Finding(
        rule_id="R-01",
        severity=Severity.ERROR,
        message=MESSAGE,
        scope=FindingScope.DOCUMENT,
        category=FindingCategory.FINANCIAL,
        subject="document",
        expected=CLAIMANT,
        actual=DESC,
    )
    return document, ReconciliationResult(status="NEEDS_REVIEW", findings=[finding])


def _assert_absent(payload: bytes) -> None:
    for name, data in _members(payload).items():
        for sentinel in SENTINELS:
            assert sentinel.encode() not in data, f"{sentinel!r} leaked into {name}"


def _all_bytes(payload: bytes) -> bytes:
    return b"\n".join(_members(payload).values()) + b"\n".join(
        name.encode() for name in _members(payload)
    )


# --------------------------------------------------------------------------
# Single-document workbook
# --------------------------------------------------------------------------


def test_redacted_workbook_leaks_no_sentinel_in_any_member():
    document, result = sentinel_document()
    payload = to_bytes(document, result, template="Full detail", redact=True)
    _assert_absent(payload)


def test_redacted_workbook_still_carries_the_decision():
    document, result = sentinel_document()
    payload = to_bytes(document, result, template="Full detail", redact=True)
    body = _all_bytes(payload)
    # The audit trail survives: the rule, the action and the claim number.
    assert b"R-01" in body
    assert b"corrected" in body
    assert b"S-1" in body


def test_source_info_states_what_was_withheld():
    document, result = sentinel_document()
    workbook = build_workbook(document, result, template="Full detail", redact=True)
    facts = {
        row[0].value: row[1].value
        for row in workbook["Source Info"].iter_rows(min_col=1, max_col=2)
        if row[0].value is not None
    }
    assert facts["Claimant data redacted"] == "yes"
    assert "claimant names" in str(facts["Withheld"])


def test_unredacted_workbook_keeps_the_sentinels():
    """The guard against over-redaction: with the toggle off, they are there."""
    document, result = sentinel_document()
    payload = to_bytes(document, result, template="Full detail", redact=False)
    body = _all_bytes(payload)
    for sentinel in (CLAIMANT, DESC, ORIG_NAME, ORIG_DESC, BEFORE, AFTER, NOTE, MESSAGE):
        assert sentinel.encode() in body, f"{sentinel!r} missing with redaction off"


def test_redacted_package_has_no_other_data_surfaces():
    """We add no comments, hyperlinks, headers/footers or dynamic sheet names."""
    from openpyxl import load_workbook

    document, result = sentinel_document()
    payload = to_bytes(document, result, template="Full detail", redact=True)
    members = _members(payload)
    assert not [n for n in members if "comment" in n.lower()]
    assert not [n for n in members if "vmlDrawing" in n]
    joined = b"\n".join(members.values())
    assert b"<hyperlink" not in joined
    assert b"<headerFooter" not in joined

    workbook = load_workbook(io.BytesIO(payload))
    assert workbook.sheetnames == [
        "Claim Detail",
        "Loss Summary",
        "Large Loss",
        "Exceptions",
        "Review History",
        "Source Info",
    ]


def test_reviewer_note_is_withheld_whole_not_scrubbed():
    document, result = sentinel_document()
    workbook = build_workbook(document, result, template="Full detail", redact=True)
    review = workbook["Review History"]
    headers = [cell.value for cell in review[1]]
    note_column = headers.index("Reviewer note") + 1
    notes = [
        review.cell(row=row, column=note_column).value
        for row in range(2, review.max_row + 1)
    ]
    assert any(note == "[withheld: claimant data redaction on]" for note in notes)


# --------------------------------------------------------------------------
# Batch archive and account workbook — the same policy
# --------------------------------------------------------------------------


def test_batch_archive_and_its_inner_workbooks_are_redacted():
    document, result = sentinel_document()
    payload, failures = build_batch_zip(
        [(document, result, None)], template="Full detail", redact=True
    )
    assert failures == []
    outer = _members(payload)
    for name, data in outer.items():
        for sentinel in SENTINELS:
            assert sentinel.encode() not in name.encode(), f"leak in entry name {name}"
            assert sentinel.encode() not in data, f"leak in {name}"
        if name.endswith(".xlsx"):
            _assert_absent(data)


def test_batch_archive_with_redaction_off_keeps_the_sentinels():
    document, result = sentinel_document()
    payload, _ = build_batch_zip(
        [(document, result, None)], template="Full detail", redact=False
    )
    outer = _members(payload)
    inner = next(data for name, data in outer.items() if name.endswith(".xlsx"))
    body = _all_bytes(inner)
    assert CLAIMANT.encode() in body
    assert DESC.encode() in body


def test_account_workbook_honours_redaction():
    document, _ = sentinel_document()
    account = build_account("Sentinel Account", [document])

    redacted = account_to_bytes(account, redact=True)
    _assert_absent(redacted)

    plain = account_to_bytes(account, redact=False)
    assert CLAIMANT.encode() in _all_bytes(plain)
